import io
import email
from email import policy
import pypdf
from docx import Document


def extract_text_from_pdf(file_bytes: bytes) -> str:
    try:
        reader = pypdf.PdfReader(io.BytesIO(file_bytes))
        pages_text = []
        for i, page in enumerate(reader.pages):
            text = page.extract_text()
            if text:
                pages_text.append(text)
        return "\n\n".join(pages_text)
    except Exception as e:
        raise ValueError(f"Error reading PDF file: {str(e)}")


def extract_text_from_docx(file_bytes: bytes) -> str:
    try:
        doc = Document(io.BytesIO(file_bytes))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    paragraphs.append(row_text)
        return "\n".join(paragraphs)
    except Exception as e:
        raise ValueError(f"Error reading DOCX file: {str(e)}")


def extract_text_from_eml(file_bytes: bytes) -> str:
    try:
        msg = email.message_from_bytes(file_bytes, policy=policy.default)
        body_parts = []
        
        # Capture email headers
        subject = msg.get("Subject", "")
        from_hdr = msg.get("From", "")
        date_hdr = msg.get("Date", "")
        to_hdr = msg.get("To", "")
        
        headers_summary = f"Date: {date_hdr}\nFrom: {from_hdr}\nTo: {to_hdr}\nSubject: {subject}\n\n"
        body_parts.append(headers_summary)

        if msg.is_multipart():
            for part in msg.walk():
                content_type = part.get_content_type()
                if content_type == "text/plain":
                    payload = part.get_payload(decode=True)
                    if payload:
                        body_parts.append(payload.decode("utf-8", errors="replace"))
        else:
            payload = msg.get_payload(decode=True)
            if payload:
                body_parts.append(payload.decode("utf-8", errors="replace"))
                
        return "".join(body_parts)
    except Exception as e:
        raise ValueError(f"Error reading EML file: {str(e)}")


def extract_text_from_txt(file_bytes: bytes) -> str:
    try:
        return file_bytes.decode("utf-8", errors="replace")
    except Exception as e:
        raise ValueError(f"Error reading TXT file: {str(e)}")


def parse_document(file_bytes: bytes, filename: str) -> str:
    lower_name = filename.lower()
    if lower_name.endswith(".pdf"):
        return extract_text_from_pdf(file_bytes)
    elif lower_name.endswith(".docx") or lower_name.endswith(".doc"):
        return extract_text_from_docx(file_bytes)
    elif lower_name.endswith(".eml") or lower_name.endswith(".msg"):
        return extract_text_from_eml(file_bytes)
    elif lower_name.endswith(".txt") or lower_name.endswith(".csv"):
        return extract_text_from_txt(file_bytes)
    else:
        # Fallback to UTF-8 decoding
        return file_bytes.decode("utf-8", errors="replace")
